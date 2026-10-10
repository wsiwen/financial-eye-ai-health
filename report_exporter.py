# -*- coding: utf-8 -*-
"""
财务慧眼 · 报告导出模块
将财务分析结果导出为 Markdown / Word 文档。
Word 版式遵循大赛《作品提交与填写说明》文案规范：A4、黑体三号大标题、
黑体四号一级标题、宋体小四正文、1.5 倍行距、首行缩进 2 字符、页脚居中页码。
"""
from __future__ import annotations

from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

import indicators as ind
from ai_analyzer import INDICATOR_CN


def build_report_markdown(company: str, code: str, indicators: dict, risks: list[dict],
                          ai_text: str, balance_check: list[dict]) -> str:
    """生成 Markdown 分析报告。"""
    periods = indicators["periods"]
    series = indicators["series"]
    latest = indicators["latest_index"]
    lines = [
        f"# {company}（{code}）财务风险分析报告",
        "",
        f"**报告期：**{periods[latest] if periods else '—'}　**生成时间：**{datetime.now():%Y-%m-%d %H:%M}",
        "",
    ]

    lines.append("## 一、报表勾稽校验")
    lines.append("")
    for item in balance_check:
        mark = "✅ 通过" if item["ok"] else "⚠️ 未通过"
        lines.append(f"- {item['period']}：{mark}；{item['detail']}")
    lines.append("")

    lines.append("## 二、核心财务指标（最新一期）")
    lines.append("")
    lines.append("| 维度 | 指标 | 数值 | 口径 |")
    lines.append("|---|---|---|---|")
    dim_order = ["偿债能力", "营运能力", "盈利能力", "现金流"]
    for dim in dim_order:
        written = False
        for key, meta in ind.INDICATOR_META.items():
            if meta[2] != dim or latest >= len(series[key]):
                continue
            v = series[key][latest]
            if v is None:
                continue
            value_str = f"{v:,.2f}" if abs(v) < 1e15 else f"{v:,.0f}"
            lines.append(f"| {dim if not written else ''} | {meta[1]} | {value_str} {meta[3]} | {meta[4]} |")
            written = True
    lines.append("")

    lines.append("## 三、风险提示")
    lines.append("")
    if risks:
        for r in risks:
            level = "🔴 高风险" if r["level"] == "high" else "🟡 预警"
            lines.append(f"- {level} **{r['name']}** {r['value']:,.2f}{r['unit']}：{r['desc']}")
    else:
        lines.append("- 未触发通用阈值风险。")
    lines.append("")

    lines.append("## 四、AI 深度解读")
    lines.append("")
    lines.append(ai_text)
    lines.append("")
    lines.append("---")
    lines.append("*本报告由「财务慧眼」AI 财报风险体检助手生成，数据来源为公开财务报表接口，仅供参考，不构成投资建议。*")
    return "\n".join(lines)


def _set_run_font(run, name_cn: str, size_pt: int, bold: bool = False):
    """设置中文字体（宋体/黑体）+ 字号。"""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name_cn)


def _add_page_number(doc: Document):
    """页脚居中插入页码域。"""
    section = doc.sections[0]
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _add_para(doc: Document, text: str, font_cn: str = "宋体", size: int = 12,
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_first: bool = True,
              bold: bool = False, spacing: float = 1.5) -> None:
    para = doc.add_paragraph()
    para.alignment = align
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if indent_first:
        pf.first_line_indent = Pt(size * 2)  # 首行缩进 2 字符
    run = para.add_run(text)
    _set_run_font(run, font_cn, size, bold)


def build_report_docx(company: str, code: str, indicators: dict, risks: list[dict],
                      ai_text: str, balance_check: list[dict], team_name: str = "外星人",
                      members: str = "乌思文", academy: str = "国际商学院") -> Document:
    """生成 Word 分析报告（符合大赛文案规范）。"""
    doc = Document()
    # 页面：A4 + 上下 2.54cm、左右 3.18cm
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.54), Cm(2.54)
    section.left_margin, section.right_margin = Cm(3.18), Cm(3.18)

    periods = indicators["periods"]
    series = indicators["series"]
    latest = indicators["latest_index"]

    # 大标题：黑体三号（16pt）居中
    _add_para(doc, f"{company}（{code}）财务风险分析报告", font_cn="黑体", size=16,
              align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=False)
    _add_para(doc, f"报告期：{periods[latest] if periods else '—'}　生成时间：{datetime.now():%Y-%m-%d}",
              size=12, align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=False)

    # 一、勾稽校验
    _add_para(doc, "一、报表勾稽校验", font_cn="黑体", size=14, indent_first=False)
    for item in balance_check:
        mark = "通过" if item["ok"] else "未通过"
        _add_para(doc, f"{item['period']}：{mark}；{item['detail']}")

    # 二、核心指标
    _add_para(doc, "二、核心财务指标（最新一期）", font_cn="黑体", size=14, indent_first=False)
    dim_order = ["偿债能力", "营运能力", "盈利能力", "现金流"]
    for dim in dim_order:
        _add_para(doc, dim, font_cn="黑体", size=12, indent_first=False)
        for key, meta in ind.INDICATOR_META.items():
            if meta[2] != dim or latest >= len(series[key]):
                continue
            v = series[key][latest]
            if v is None:
                continue
            value_str = f"{v:,.2f}" if abs(v) < 1e15 else f"{v:,.0f}"
            _add_para(doc, f"{meta[1]}：{value_str} {meta[3]}（{meta[4]}）")

    # 三、风险提示
    _add_para(doc, "三、风险提示", font_cn="黑体", size=14, indent_first=False)
    if risks:
        for r in risks:
            level = "高风险" if r["level"] == "high" else "预警"
            _add_para(doc, f"【{level}】{r['name']} {r['value']:,.2f}{r['unit']}：{r['desc']}")
    else:
        _add_para(doc, "未触发通用阈值风险。")

    # 四、AI 解读（把 markdown 简单转成段落）
    _add_para(doc, "四、AI 深度解读", font_cn="黑体", size=14, indent_first=False)
    for line in (ai_text or "").splitlines():
        clean = line.strip()
        if not clean:
            continue
        if clean.startswith("### "):
            _add_para(doc, clean[4:], font_cn="黑体", size=12, indent_first=False)
        elif clean.startswith("## "):
            _add_para(doc, clean[3:], font_cn="黑体", size=12, indent_first=False)
        elif clean.startswith("- "):
            _add_para(doc, clean[2:], indent_first=False)
        elif clean.startswith("> "):
            _add_para(doc, clean[2:])
        else:
            _add_para(doc, clean)

    # 落款（右侧）：学院、团队、成员、日期
    _add_para(doc, "", indent_first=False)
    for text in [f"推荐学院：{academy}", f"团队名称：{team_name}", f"成员：{members}",
                 f"日期：{datetime.now():%Y年%m月%d日}"]:
        _add_para(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT, indent_first=False)

    _add_page_number(doc)
    return doc
